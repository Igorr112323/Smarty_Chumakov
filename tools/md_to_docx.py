#!/usr/bin/env python3
"""Конвертер Markdown → DOCX (для сопроводительных документов).

Поддерживает: заголовки, абзацы, списки, цитаты, таблицы, встроенный код,
горизонтальные линии. Форматирование намеренно простое и предсказуемое.

    python tools/md_to_docx.py Промты_для_ИИ.md --out Промты_для_ИИ.docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def add_inline(paragraph, text: str) -> None:
    """Разобрать **жирный**, *курсив*, `код`, экранирование."""
    token_re = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
    for chunk in token_re.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x22, 0x44, 0x88)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        else:
            paragraph.add_run(chunk.replace("\\", ""))


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    i = 0
    in_table = False
    table_buffer: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buffer, in_table
        if not table_buffer:
            in_table = False
            return
        rows = [r for r in table_buffer if not re.match(r"^\s*\|[\s\-:|]+\|\s*$", r[0])] if False else table_buffer
        header = rows[0]
        body = rows[1:]
        table = doc.add_table(rows=1, cols=len(header))
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, header):
            cell.paragraphs[0].text = ""
            add_inline(cell.paragraphs[0], text)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in body:
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                cell.paragraphs[0].text = ""
                add_inline(cell.paragraphs[0], text)
        doc.add_paragraph()
        table_buffer = []
        in_table = False

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r"-{2,}|:?-+:?", c) for c in cells) and table_buffer:
                i += 1
                continue
            table_buffer.append(cells)
            in_table = True
            i += 1
            continue

        if in_table:
            flush_table()

        if not line.strip():
            i += 1
            continue

        if line.startswith("---"):
            doc.add_paragraph("─" * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(4, len(heading.group(1)))
            text = heading.group(2)
            paragraph = doc.add_heading("", level=level)
            add_inline(paragraph, text)
            i += 1
            continue

        if line.startswith("> "):
            paragraph = doc.add_paragraph(style="Intense Quote")
            add_inline(paragraph, line[2:])
            i += 1
            continue

        checkbox = re.match(r"^\s*-\s+\[( |x)\]\s+(.*)$", line)
        if checkbox:
            mark = "☒" if checkbox.group(1) == "x" else "☐"
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, f"{mark} {checkbox.group(2)}")
            i += 1
            continue

        bullet = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(2))
            i += 1
            continue

        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            i += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, line)
        i += 1

    if in_table:
        flush_table()

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"{md_path} → {docx_path}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Markdown → DOCX")
    parser.add_argument("source", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    convert(args.source, args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
